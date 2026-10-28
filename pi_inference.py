import json
import os
import queue
import re
import shutil
import sqlite3
import sys
import tempfile
import time
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from threading import Event, Lock, Thread
from typing import Dict, List, Optional, Tuple

import cv2
import numpy as np
import torch
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from transformers import AutoConfig, AutoImageProcessor, AutoModelForImageClassification
from ultralytics import YOLO


BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = Path(os.getenv("MIVOLO_CONFIG_FILE", BASE_DIR / "config.json"))
try:
    with CONFIG_FILE.open("r", encoding="utf-8") as config_file:
        app_config = json.load(config_file)
    if not isinstance(app_config, dict) or not isinstance(app_config.get("pi_inference"), dict):
        raise ValueError("config.json must contain a pi_inference object")
    server_config = app_config["pi_inference"]
except (OSError, json.JSONDecodeError, ValueError) as exc:
    raise RuntimeError(f"Unable to load MiVOLO server config {CONFIG_FILE}: {exc}") from exc


def _config_value(name: str, expected_type: type, default):
    value = server_config.get(name, default)
    if type(value) is not expected_type:
        raise RuntimeError(f"pi_inference.{name} in {CONFIG_FILE} must be a {expected_type.__name__}")
    return value


MODEL_ID = _config_value("model_id", str, "iitolstykh/mivolo_v2")
RESULTS_FILE = Path(_config_value("results_file", str, "predictions.xlsx"))
IMAGES_DIR = Path(_config_value("images_dir", str, "prediction_images"))
NUM_THREADS = _config_value("num_threads", int, 2)
MAX_QUEUE_SIZE = _config_value("max_queue_size", int, 15)
RETENTION_DAYS = _config_value("retention_days", int, 7)
RETENTION_CHECK_INTERVAL = _config_value("retention_check_interval_seconds", int, 3600)
MAINTENANCE_DB = Path(_config_value("maintenance_db", str, "mivolo_maintenance.sqlite3"))
face_detector_config = server_config.get("face_detector", {})
if not isinstance(face_detector_config, dict):
    raise RuntimeError(f"pi_inference.face_detector in {CONFIG_FILE} must be an object")
FACE_MODEL_PATH = Path(face_detector_config.get("model_path", "yolov11n-face.pt"))
FACE_CONFIDENCE = face_detector_config.get("confidence", 0.25)
FACE_IMAGE_SIZE = face_detector_config.get("image_size", 640)
if type(FACE_CONFIDENCE) not in (int, float) or not 0 < FACE_CONFIDENCE <= 1:
    raise RuntimeError("pi_inference.face_detector.confidence must be in (0, 1]")
if type(FACE_IMAGE_SIZE) is not int or FACE_IMAGE_SIZE < 1:
    raise RuntimeError("pi_inference.face_detector.image_size must be a positive integer")
if (
    not MODEL_ID
    or NUM_THREADS < 1
    or MAX_QUEUE_SIZE < 1
    or RETENTION_DAYS < 1
    or RETENTION_CHECK_INTERVAL < 1
):
    raise RuntimeError("model_id, thread/queue sizes, and retention settings must be positive")
if not RESULTS_FILE.is_absolute():
    RESULTS_FILE = BASE_DIR / RESULTS_FILE
if not IMAGES_DIR.is_absolute():
    IMAGES_DIR = BASE_DIR / IMAGES_DIR
if not MAINTENANCE_DB.is_absolute():
    MAINTENANCE_DB = BASE_DIR / MAINTENANCE_DB
if not FACE_MODEL_PATH.is_absolute():
    FACE_MODEL_PATH = BASE_DIR / FACE_MODEL_PATH
BASE_HEADERS = [
    "timestamp_utc",
    "body_bbox",
    "face_detected",
    "face_bbox",
    "age",
    "gender",
    "gender_probability",
]
ADDITIONAL_HEADERS = [
    "track_id",
    "request_id",
    "frame_image",
    "body_image",
    "face_image",
    "inference_seconds",
]
HEADERS = BASE_HEADERS + ADDITIONAL_HEADERS
OBSOLETE_HEADERS = {
    "image_filename",
    "frame_image_path",
    "body_image_path",
    "face_image_path",
}
LEGACY_HEADERS = [header for header in HEADERS if header != "inference_seconds"]
IMAGE_COLUMNS = {"frame": "frame_image", "body": "body_image", "face": "face_image"}

model = None
model_config = None
image_processor = None
inference_lock = Lock()
excel_lock = Lock()
prediction_queue = queue.Queue(maxsize=MAX_QUEUE_SIZE)
prediction_worker = None
retention_stop = Event()
retention_worker = None


class PredictionJob:
    def __init__(self, image_bytes: bytes, filename: str, body_bbox: str, track_id: str):
        self.image_bytes = image_bytes
        self.filename = filename
        self.body_bbox = body_bbox
        self.track_id = track_id
        self.done = Event()
        self.result = None
        self.error = None


@asynccontextmanager
async def lifespan(_: FastAPI):
    global model, model_config, image_processor, prediction_worker, retention_worker
    global face_yolo_model

    torch.set_num_threads(NUM_THREADS)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    dtype = torch.float16 if device.type == "cuda" else torch.float32

    model_config = AutoConfig.from_pretrained(MODEL_ID, trust_remote_code=True)
    model = AutoModelForImageClassification.from_pretrained(
        MODEL_ID,
        trust_remote_code=True,
        torch_dtype=dtype,
    ).to(device)
    model.eval()
    image_processor = AutoImageProcessor.from_pretrained(MODEL_ID, trust_remote_code=True)
    if FACE_MODEL_PATH.is_file():
        try:
            face_yolo_model = YOLO(str(FACE_MODEL_PATH))
            print(f"Loaded YOLO face detector: {FACE_MODEL_PATH}", flush=True)
        except (OSError, RuntimeError, ValueError) as exc:
            face_yolo_model = None
            print(
                f"Unable to load YOLO face detector {FACE_MODEL_PATH}: {exc}; "
                "using Haar cascade fallback.",
                file=sys.stderr,
                flush=True,
            )
    else:
        face_yolo_model = None
        print(
            f"YOLO face detector weights not found at {FACE_MODEL_PATH}; "
            "using Haar cascade fallback.",
            file=sys.stderr,
            flush=True,
        )
    _migrate_legacy_workbook()
    _migrate_legacy_images()
    prediction_worker = Thread(
        target=_prediction_worker,
        daemon=True,
        name="mivolo-inference-worker",
    )
    prediction_worker.start()
    retention_stop.clear()
    retention_worker = Thread(
        target=_retention_worker,
        daemon=True,
        name="mivolo-retention-worker",
    )
    retention_worker.start()
    try:
        yield
    finally:
        retention_stop.set()
        retention_worker.join()
        prediction_queue.put(None)
        prediction_worker.join()


app = FastAPI(title="MiVOLO Age and Gender API", lifespan=lifespan)
face_yolo_model = None
face_detector = None
if hasattr(cv2, "CascadeClassifier"):
    candidate_detector = cv2.CascadeClassifier(
        os.path.join(cv2.data.haarcascades, "haarcascade_frontalface_default.xml")
    )
    if not candidate_detector.empty():
        face_detector = candidate_detector
    else:
        print(
            "Warning: OpenCV's frontal face cascade could not be loaded; "
            "age-gender inference will use the body-only fallback.",
            flush=True,
        )
else:
    print(
        "Warning: this OpenCV build has no CascadeClassifier; "
        "age-gender inference will use the body-only fallback.",
        flush=True,
    )


def _parse_bbox(value: str, image_width: int, image_height: int) -> Tuple[int, int, int, int]:
    try:
        bbox = json.loads(value)
        if (
            not isinstance(bbox, list)
            or len(bbox) != 4
            or any(type(coord) is not int for coord in bbox)
        ):
            raise ValueError
        x1, y1, x2, y2 = bbox
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(
            status_code=422,
            detail="body_bbox must be a JSON array of four integers",
        ) from exc

    if x1 < 0 or y1 < 0 or x2 > image_width or y2 > image_height or x1 >= x2 or y1 >= y2:
        raise HTTPException(
            status_code=422,
            detail=(
                f"body_bbox must be a non-empty [x1, y1, x2, y2] box "
                f"within {image_width}x{image_height}"
            ),
        )
    return x1, y1, x2, y2


def _validate_track_id(track_id: str) -> str:
    track_id = track_id.strip()
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,99}", track_id):
        raise HTTPException(
            status_code=422,
            detail=(
                "track_id must be 1-100 characters, start with a letter or digit, "
                "and contain only letters, digits, '.', '_' or '-'."
            ),
        )
    return track_id


def _detect_face(
    image: np.ndarray,
    body_bbox: Tuple[int, int, int, int],
) -> Tuple[Optional[np.ndarray], Optional[List[int]]]:
    x1, y1, x2, y2 = body_bbox
    body_region = image[y1:y2, x1:x2]
    if face_yolo_model is not None:
        try:
            results = face_yolo_model.predict(
                source=body_region,
                conf=FACE_CONFIDENCE,
                imgsz=FACE_IMAGE_SIZE,
                verbose=False,
                device="cpu",
            )
            if results and results[0].boxes is not None and len(results[0].boxes):
                boxes = results[0].boxes.xyxy.cpu().numpy()
                face_x1, face_y1, face_x2, face_y2 = (
                    int(round(value))
                    for value in max(
                        boxes,
                        key=lambda box: max(0, box[2] - box[0]) * max(0, box[3] - box[1]),
                    )
                )
                face_x1 = max(0, min(body_region.shape[1], face_x1))
                face_x2 = max(0, min(body_region.shape[1], face_x2))
                face_y1 = max(0, min(body_region.shape[0], face_y1))
                face_y2 = max(0, min(body_region.shape[0], face_y2))
                if face_x1 < face_x2 and face_y1 < face_y2:
                    abs_box = [x1 + face_x1, y1 + face_y1, x1 + face_x2, y1 + face_y2]
                    return image[abs_box[1]:abs_box[3], abs_box[0]:abs_box[2]], abs_box
        except Exception as exc:
            print(
                f"YOLO face detection failed; trying Haar cascade: {exc}",
                file=sys.stderr,
                flush=True,
            )

    if face_detector is None:
        return None, None

    gray = cv2.cvtColor(body_region, cv2.COLOR_BGR2GRAY)
    faces = face_detector.detectMultiScale(
        gray,
        scaleFactor=1.1,
        minNeighbors=5,
        minSize=(24, 24),
    )
    if len(faces) == 0:
        return None, None

    face_x, face_y, face_width, face_height = (
        int(value)
        for value in max(faces, key=lambda box: box[2] * box[3])
    )
    abs_x1, abs_y1 = x1 + face_x, y1 + face_y
    abs_x2, abs_y2 = abs_x1 + face_width, abs_y1 + face_height
    return image[abs_y1:abs_y2, abs_x1:abs_x2], [abs_x1, abs_y1, abs_x2, abs_y2]


def _prepare_input(crop: np.ndarray) -> torch.Tensor:
    crop_rgb = cv2.cvtColor(crop, cv2.COLOR_BGR2RGB)
    inputs = image_processor(images=[crop_rgb])["pixel_values"]
    return inputs.to(device=model.device, dtype=model.dtype)


def _encode_jpeg(image: np.ndarray) -> bytes:
    success, encoded = cv2.imencode(".jpg", image)
    if not success:
        raise RuntimeError("Failed to encode an image as JPEG")
    return encoded.tobytes()


def _daily_results_file(local_date: str) -> Path:
    return RESULTS_FILE.parent / local_date / RESULTS_FILE.name


def _daily_images_dir(local_date: str) -> Path:
    return IMAGES_DIR / local_date


def _save_request_images(
    local_date: str,
    track_id: str,
    request_id: str,
    frame: np.ndarray,
    body_crop: np.ndarray,
    face_crop: Optional[np.ndarray],
) -> Dict[str, bytes]:
    request_dir = _daily_images_dir(local_date) / track_id / request_id
    request_dir.mkdir(parents=True, exist_ok=True)

    images = {"frame": frame, "body": body_crop}
    if face_crop is not None:
        images["face"] = face_crop

    saved_images = {}
    for image_type, crop in images.items():
        image_bytes = _encode_jpeg(crop)
        image_path = request_dir / f"{image_type}.jpg"
        image_path.write_bytes(image_bytes)
        saved_images[image_type] = image_bytes
    return saved_images


def _thumbnail(image_bytes: bytes) -> ExcelImage:
    image = ExcelImage(BytesIO(image_bytes))
    scale = min(120 / image.width, 120 / image.height, 1)
    image.width = int(image.width * scale)
    image.height = int(image.height * scale)
    return image


def _remove_obsolete_columns(worksheet) -> None:
    headers = [cell.value for cell in worksheet[1]]
    obsolete_columns = [
        index
        for index, header in enumerate(headers)
        if header in OBSOLETE_HEADERS
    ]
    if not obsolete_columns:
        return

    def adjusted_column(column: int) -> int:
        return column - sum(removed <= column for removed in obsolete_columns)

    for image in worksheet._images:
        anchor = image.anchor
        if not isinstance(anchor, str):
            anchor._from.col = adjusted_column(anchor._from.col)
            if hasattr(anchor, "to"):
                anchor.to.col = adjusted_column(anchor.to.col)

    for column in reversed(obsolete_columns):
        worksheet.delete_cols(column + 1)


def _save_prediction(
    row: Dict[str, object],
    embedded_images: Dict[str, bytes],
    results_file: Path,
) -> None:
    with excel_lock:
        if results_file.exists():
            workbook = load_workbook(results_file)
            if "Predictions" in workbook.sheetnames:
                worksheet = workbook["Predictions"]
            else:
                worksheet = workbook.create_sheet("Predictions")
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Predictions"

        try:
            _remove_obsolete_columns(worksheet)
            if worksheet.max_row == 1 and worksheet.cell(row=1, column=1).value is None:
                for column, header in enumerate(HEADERS, start=1):
                    worksheet.cell(row=1, column=column, value=header)
            else:
                existing_headers = [cell.value for cell in worksheet[1]]
                if existing_headers == BASE_HEADERS:
                    for column, header in enumerate(ADDITIONAL_HEADERS, start=len(BASE_HEADERS) + 1):
                        worksheet.cell(row=1, column=column, value=header)
                elif existing_headers == LEGACY_HEADERS:
                    worksheet.cell(row=1, column=len(HEADERS), value="inference_seconds")
                elif existing_headers != HEADERS:
                    raise ValueError(
                        f"{results_file} has an incompatible Predictions sheet header"
                    )

            worksheet.append([row.get(header) for header in HEADERS])
            row_number = worksheet.max_row
            image_column_indices = {
                name: HEADERS.index(column) + 1
                for name, column in IMAGE_COLUMNS.items()
            }
            for image_type, image_bytes in embedded_images.items():
                image = _thumbnail(image_bytes)
                column_index = image_column_indices[image_type]
                image.anchor = worksheet.cell(row=row_number, column=column_index).coordinate
                worksheet.add_image(image)
                worksheet.column_dimensions[
                    worksheet.cell(row=1, column=column_index).column_letter
                ].width = 18
            if embedded_images:
                worksheet.row_dimensions[row_number].height = 70

            results_file.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(results_file)
        finally:
            workbook.close()


def _local_datetime(value, fallback: datetime) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str) and value:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    else:
        return fallback
    return parsed.astimezone() if parsed.tzinfo else parsed.astimezone()


def _migrate_legacy_workbook() -> None:
    if not RESULTS_FILE.exists():
        return

    fallback = datetime.fromtimestamp(RESULTS_FILE.stat().st_mtime).astimezone()
    workbook = load_workbook(RESULTS_FILE)
    records_by_date = {}
    try:
        if "Predictions" not in workbook.sheetnames:
            return
        worksheet = workbook["Predictions"]
        _remove_obsolete_columns(worksheet)
        headers = [cell.value for cell in worksheet[1]]
        if not headers or "timestamp_utc" not in headers:
            raise ValueError(f"{RESULTS_FILE} has no timestamp_utc column")
        timestamp_column = headers.index("timestamp_utc") + 1
        if headers != HEADERS and headers != LEGACY_HEADERS and headers != BASE_HEADERS:
            raise ValueError(f"{RESULTS_FILE} has an incompatible Predictions sheet header")

        images_by_row = {}
        image_header_columns = {
            headers.index(column): image_type
            for image_type, column in IMAGE_COLUMNS.items()
            if column in headers
        }
        for image in worksheet._images:
            anchor = image.anchor
            row_index = anchor._from.row
            image_type = image_header_columns.get(anchor._from.col)
            if image_type:
                images_by_row.setdefault(row_index, {})[image_type] = image._data()

        for row_index in range(1, worksheet.max_row):
            excel_row = row_index + 1
            values = {
                header: worksheet.cell(row=excel_row, column=column_index + 1).value
                for column_index, header in enumerate(headers)
                if header in HEADERS
            }
            local_date = _local_datetime(
                worksheet.cell(row=excel_row, column=timestamp_column).value,
                fallback,
            ).date().isoformat()
            records_by_date.setdefault(local_date, []).append(
                (values, images_by_row.get(row_index, {}))
            )
    finally:
        workbook.close()

    if not records_by_date:
        records_by_date[fallback.date().isoformat()] = []

    RESULTS_FILE.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".mivolo-migration-",
        dir=RESULTS_FILE.parent,
    ) as staging_dir:
        staged_files = []
        for local_date, records in records_by_date.items():
            destination = _daily_results_file(local_date)
            if not records and destination.exists():
                continue
            if not records and not destination.exists():
                continue
            staged_file = Path(staging_dir) / local_date / RESULTS_FILE.name
            staged_file.parent.mkdir(parents=True, exist_ok=True)
            existing_request_ids = set()
            if destination.exists():
                shutil.copy2(destination, staged_file)
                existing = load_workbook(staged_file, read_only=True, data_only=True)
                try:
                    sheet = existing["Predictions"]
                    existing_headers = [cell.value for cell in sheet[1]]
                    if "request_id" in existing_headers:
                        request_column = existing_headers.index("request_id") + 1
                        existing_request_ids = {
                            sheet.cell(row=row, column=request_column).value
                            for row in range(2, sheet.max_row + 1)
                        }
                finally:
                    existing.close()

            for record, embedded_images in records:
                request_id = record.get("request_id")
                if request_id and request_id in existing_request_ids:
                    continue
                _save_prediction(record, embedded_images, staged_file)
                if request_id:
                    existing_request_ids.add(request_id)
            staged_files.append((staged_file, destination))

        for staged_file, destination in staged_files:
            if not staged_file.exists():
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            os.replace(staged_file, destination)
    RESULTS_FILE.unlink()
    print(
        f"Migrated legacy workbook into {len(records_by_date)} local-date workbook(s)",
        flush=True,
    )


def _migrate_legacy_images() -> None:
    if not IMAGES_DIR.exists():
        return

    for track_dir in IMAGES_DIR.iterdir():
        if not track_dir.is_dir() or re.fullmatch(r"\d{4}-\d{2}-\d{2}", track_dir.name):
            continue
        for request_dir in list(track_dir.iterdir()):
            if not request_dir.is_dir():
                continue
            files = [path for path in request_dir.iterdir() if path.is_file()]
            if not files:
                continue
            local_date = datetime.fromtimestamp(max(path.stat().st_mtime for path in files)).astimezone().date().isoformat()
            destination = _daily_images_dir(local_date) / track_dir.name / request_dir.name
            if destination.exists():
                raise FileExistsError(f"Cannot migrate images; destination already exists: {destination}")
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(request_dir), str(destination))
        if not any(track_dir.iterdir()):
            track_dir.rmdir()


def _prune_workbook(path: Path, cutoff: datetime) -> int:
    workbook = load_workbook(path)
    try:
        if "Predictions" not in workbook.sheetnames:
            raise ValueError(f"{path} has no Predictions sheet")
        worksheet = workbook["Predictions"]
        headers = [cell.value for cell in worksheet[1]]
        if "timestamp_utc" not in headers:
            raise ValueError(f"{path} has no timestamp_utc column")
        timestamp_column = headers.index("timestamp_utc") + 1
        fallback = datetime.fromtimestamp(path.stat().st_mtime).astimezone()
        deleted_rows = {
            row
            for row in range(2, worksheet.max_row + 1)
            if _local_datetime(
                worksheet.cell(row=row, column=timestamp_column).value,
                fallback,
            ) < cutoff
        }
        if not deleted_rows:
            return 0

        kept_images = []
        for image in worksheet._images:
            anchor_row = image.anchor._from.row + 1
            if anchor_row in deleted_rows:
                continue
            image.anchor._from.row -= sum(row < anchor_row for row in deleted_rows)
            if hasattr(image.anchor, "to"):
                image.anchor.to.row -= sum(row < anchor_row for row in deleted_rows)
            kept_images.append(image)
        worksheet._images = kept_images

        for row in sorted(deleted_rows, reverse=True):
            worksheet.delete_rows(row)
        workbook.save(path)
        return len(deleted_rows)
    finally:
        workbook.close()


def _prune_old_images(cutoff: datetime) -> int:
    if not IMAGES_DIR.exists():
        return 0
    cutoff_timestamp = cutoff.timestamp()
    removed = 0
    for request_dir in sorted(IMAGES_DIR.rglob("*"), key=lambda path: len(path.parts), reverse=True):
        if not request_dir.is_dir():
            continue
        children = list(request_dir.iterdir())
        if children and all(child.is_file() for child in children):
            if max(child.stat().st_mtime for child in children) < cutoff_timestamp:
                shutil.rmtree(request_dir)
                removed += 1
        elif not children:
            request_dir.rmdir()
    return removed


def _record_retention_check(
    checked_at: datetime,
    cutoff: datetime,
    removed_image_requests: int,
    removed_excel_rows: int,
) -> None:
    MAINTENANCE_DB.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(MAINTENANCE_DB) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS retention_checks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                checked_at_local TEXT NOT NULL,
                cutoff_local TEXT NOT NULL,
                removed_image_requests INTEGER NOT NULL,
                removed_excel_rows INTEGER NOT NULL
            )
            """
        )
        connection.execute(
            """
            INSERT INTO retention_checks
                (checked_at_local, cutoff_local, removed_image_requests, removed_excel_rows)
            VALUES (?, ?, ?, ?)
            """,
            (
                checked_at.isoformat(),
                cutoff.isoformat(),
                removed_image_requests,
                removed_excel_rows,
            ),
        )


def _run_retention_check() -> None:
    checked_at = datetime.now().astimezone()
    cutoff = checked_at - timedelta(days=RETENTION_DAYS)
    removed_image_requests = _prune_old_images(cutoff)
    removed_excel_rows = 0

    candidates = []
    if RESULTS_FILE.exists():
        candidates.append(RESULTS_FILE)
    if RESULTS_FILE.parent.exists():
        for date_dir in RESULTS_FILE.parent.iterdir():
            if date_dir.is_dir() and re.fullmatch(r"\d{4}-\d{2}-\d{2}", date_dir.name):
                daily_file = date_dir / RESULTS_FILE.name
                if daily_file.exists():
                    candidates.append(daily_file)

    for workbook_path in candidates:
        with excel_lock:
            removed_excel_rows += _prune_workbook(workbook_path, cutoff)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                worksheet = workbook["Predictions"]
                has_records = worksheet.max_row > 1
            finally:
                workbook.close()
        if not has_records:
            workbook_path.unlink()
            if workbook_path.parent != RESULTS_FILE.parent and not any(workbook_path.parent.iterdir()):
                workbook_path.parent.rmdir()

    _record_retention_check(
        checked_at,
        cutoff,
        removed_image_requests,
        removed_excel_rows,
    )
    print(
        f"Retention check complete at {checked_at.isoformat()}: "
        f"removed {removed_image_requests} image requests and "
        f"{removed_excel_rows} Excel rows older than {RETENTION_DAYS} days.",
        flush=True,
    )


def _retention_worker() -> None:
    while not retention_stop.is_set():
        try:
            _run_retention_check()
        except Exception as exc:
            print(f"Retention check failed: {exc}", file=sys.stderr, flush=True)
        if retention_stop.wait(RETENTION_CHECK_INTERVAL):
            break


@app.post("/predict")
def predict(
    image: UploadFile = File(..., description="Full-size image frame"),
    body_bbox: str = Form(..., description="Body box as JSON: [x1, y1, x2, y2]"),
    track_id: str = Form(..., description="Tracking/person ID used to group saved images"),
):
    received_at = datetime.now(timezone.utc)
    print(
        f"[{received_at.isoformat()}] Request received: track_id={track_id!r}, "
        f"filename={image.filename!r}, body_bbox={body_bbox}",
        flush=True,
    )
    track_id = _validate_track_id(track_id)

    try:
        job = PredictionJob(
            image_bytes=image.file.read(),
            filename=image.filename or "",
            body_bbox=body_bbox,
            track_id=track_id,
        )
    finally:
        image.file.close()

    try:
        prediction_queue.put_nowait(job)
    except queue.Full as exc:
        raise HTTPException(
            status_code=503,
            detail="Inference queue is full; retry the request shortly",
            headers={"Retry-After": "2"},
        ) from exc

    job.done.wait()
    if job.error is not None:
        if isinstance(job.error, HTTPException):
            raise job.error
        print(f"Unhandled prediction error for track_id={track_id!r}: {job.error}", flush=True)
        raise HTTPException(status_code=500, detail="Prediction processing failed") from job.error
    return job.result


def _process_prediction(job: PredictionJob):
    if model is None or model_config is None or image_processor is None:
        raise HTTPException(status_code=503, detail="The age-gender model is not initialized")

    encoded_image = np.frombuffer(job.image_bytes, dtype=np.uint8)
    frame = cv2.imdecode(encoded_image, cv2.IMREAD_COLOR)
    if frame is None:
        raise HTTPException(status_code=400, detail="Uploaded file is not a valid image")

    height, width = frame.shape[:2]
    parsed_body_bbox = _parse_bbox(job.body_bbox, width, height)
    x1, y1, x2, y2 = parsed_body_bbox
    body_crop = frame[y1:y2, x1:x2]
    face_crop, face_bbox = _detect_face(frame, parsed_body_bbox)

    analysis_start = datetime.now(timezone.utc)
    request_id = (
        f"{analysis_start.strftime('%Y%m%dT%H%M%S_%fZ')}_{uuid.uuid4().hex[:8]}"
    )
    try:
        body_input = _prepare_input(body_crop)
        face_model_crop = face_crop if face_crop is not None else np.zeros_like(body_crop)
        faces_input = _prepare_input(face_model_crop)
        inference_start = time.perf_counter()
        with inference_lock, torch.inference_mode():
            output = model(faces_input=faces_input, body_input=body_input)
        inference_seconds = time.perf_counter() - inference_start

        age = output.age_output[0].item()
        gender_index = output.gender_class_idx[0].item()
        gender_labels = model_config.gender_id2label
        gender = gender_labels.get(gender_index, gender_labels.get(str(gender_index)))
        if gender is None:
            raise ValueError(f"No gender label configured for class index {gender_index}")
        gender_probability = output.gender_probs[0].item()
    except HTTPException:
        raise
    except (OSError, RuntimeError, ValueError, KeyError) as exc:
        print(f"Inference failed for track_id={job.track_id!r}: {exc}", flush=True)
        raise HTTPException(status_code=500, detail="Age-gender inference failed") from exc

    processed_at = datetime.now().astimezone()
    local_date = processed_at.date().isoformat()
    try:
        saved_images = _save_request_images(
            local_date,
            job.track_id,
            request_id,
            frame,
            body_crop,
            face_crop,
        )
        request_dir = (
            _daily_images_dir(local_date) / job.track_id / request_id
        ).resolve()
        saved_paths = {
            image_type: str(request_dir / f"{image_type}.jpg")
            for image_type in saved_images
        }
    except (OSError, RuntimeError) as exc:
        print(f"Failed to save request images for track_id={job.track_id!r}: {exc}", flush=True)
        raise HTTPException(status_code=500, detail="Failed to save request images") from exc

    result = {
        "age": round(age, 2),
        "gender": gender,
        "gender_probability": gender_probability,
        "face_detected": face_crop is not None,
        "face_bbox": face_bbox,
        "body_bbox": parsed_body_bbox,
        "track_id": job.track_id,
        "request_id": request_id,
        "inference_seconds": round(inference_seconds, 6),
        "saved_images": saved_paths,
    }
    record = {
        "timestamp_utc": processed_at.astimezone(timezone.utc).isoformat(),
        "body_bbox": json.dumps(parsed_body_bbox),
        "face_detected": result["face_detected"],
        "face_bbox": json.dumps(face_bbox) if face_bbox is not None else "",
        "age": result["age"],
        "gender": result["gender"],
        "gender_probability": result["gender_probability"],
        "track_id": job.track_id,
        "request_id": request_id,
        "frame_image": None,
        "body_image": None,
        "face_image": None,
        "inference_seconds": result["inference_seconds"],
    }
    results_file = _daily_results_file(local_date)
    try:
        _save_prediction(record, saved_images, results_file)
    except (OSError, ValueError, KeyError) as exc:
        print(f"Failed to save prediction to {results_file}: {exc}", flush=True)
        raise HTTPException(
            status_code=500,
            detail="Failed to save prediction to Excel workbook",
        ) from exc

    print(
        f"[{record['timestamp_utc']}] Processing complete: "
        f"track_id={job.track_id!r}, filename={job.filename!r}, "
        f"result={json.dumps(result)} excel={str(results_file.resolve())!r}",
        flush=True,
    )
    return result


def _prediction_worker() -> None:
    while True:
        job = prediction_queue.get()
        try:
            if job is None:
                return
            try:
                job.result = _process_prediction(job)
            except Exception as exc:
                job.error = exc
        finally:
            if job is not None:
                job.done.set()
            prediction_queue.task_done()
